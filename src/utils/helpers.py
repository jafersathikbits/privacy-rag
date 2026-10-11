# src/utils/helpers.py
import hashlib
import os
import pickle
import shutil
import numpy as np

from src.config import (
    CHUNK_SIZE, CHUNK_OVERLAP, BATCH_SIZE,
    EMBEDDING_MODEL, QUERY_INSTRUCTION,
)

# Heavy libraries (torch, sentence-transformers, PyMuPDF, docx, tesseract) are
# imported lazily so that importing this module is fast and only the
# extractors actually needed get loaded.

# ------------------------------------------------------------------
# Text extraction
# ------------------------------------------------------------------
IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.bmp', '.tiff'}


def extract_text_from_file(file_path):
    """Extract text from PDF, DOCX, TXT/MD, or image files."""
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext in ('.txt', '.md'):
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        elif ext == '.pdf':
            import fitz  # PyMuPDF
            with fitz.open(file_path) as doc:
                return "\n".join(page.get_text() for page in doc)
        elif ext == '.docx':
            import docx
            doc = docx.Document(file_path)
            parts = [para.text for para in doc.paragraphs]
            # Tables were previously ignored entirely.
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        elif ext in IMAGE_EXTENSIONS:
            import pytesseract
            from PIL import Image
            tesseract_cmd = shutil.which('tesseract')
            if tesseract_cmd:
                pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
            elif os.name == 'nt':
                # Common Windows install path fallback
                pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
            with Image.open(file_path) as image:
                return pytesseract.image_to_string(image)
        else:
            return ""
    except Exception as e:
        print(f"Error extracting {file_path}: {e}")
        return ""

# ------------------------------------------------------------------
# Chunking
# ------------------------------------------------------------------
def chunk_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """Split text into overlapping chunks of words."""
    words = text.split()
    if not words:
        return []
    if overlap >= chunk_size:
        raise ValueError("CHUNK_OVERLAP must be smaller than CHUNK_SIZE")
    chunks = []
    step = chunk_size - overlap
    for i in range(0, len(words), step):
        chunks.append(' '.join(words[i:i + chunk_size]))
        # Stop once the window reaches the end; the old loop emitted extra
        # tail chunks that were pure subsets of the previous chunk.
        if i + chunk_size >= len(words):
            break
    return chunks

# ------------------------------------------------------------------
# Embedding (cached model with GPU support)
# ------------------------------------------------------------------
_model = None


def get_embedding_model(model_name=EMBEDDING_MODEL):
    """Load and cache the embedding model, preferring GPU if available."""
    global _model
    if _model is None:
        import torch
        from sentence_transformers import SentenceTransformer
        device = "cuda" if torch.cuda.is_available() else "cpu"
        _model = SentenceTransformer(model_name, device=device)
    return _model


def embed_texts(texts, batch_size=BATCH_SIZE, show_progress_bar=True):
    """Embed documents/passages (no instruction prefix). Returns unit vectors."""
    model = get_embedding_model()
    return model.encode(
        texts,
        batch_size=batch_size,
        show_progress_bar=show_progress_bar,
        convert_to_numpy=True,
        normalize_embeddings=True,
    ).astype(np.float32)


def embed_query(query, model=None):
    """Embed a single user query, applying QUERY_INSTRUCTION if configured.
    BGE-style models require a prefix on queries, but NOT on documents."""
    if model is None:
        model = get_embedding_model()
    text = f"{QUERY_INSTRUCTION}{query}" if QUERY_INSTRUCTION else query
    return model.encode([text], convert_to_numpy=True, normalize_embeddings=True)[0].astype(np.float32)

# ------------------------------------------------------------------
# Orthogonal Masking (random permutation + sign flips)
# ------------------------------------------------------------------
def generate_transform(dim):
    """Generate a random permutation and sign flip vector for masking."""
    rng = np.random.default_rng()
    perm = rng.permutation(dim)
    signs = rng.choice(np.array([-1, 1], dtype=np.float32), size=dim)
    return perm, signs


def apply_transform(embedding, perm, signs):
    """Apply the transform (permutation + sign flip) to an embedding vector."""
    signs = signs.astype(np.float32, copy=False)
    if embedding.ndim == 1:
        return signs * embedding[perm]
    return signs * embedding[:, perm]

# ------------------------------------------------------------------
# Merkle Mountain Range (MMR)
# ------------------------------------------------------------------
def _hash_leaf(leaf):
    return hashlib.sha256(leaf.encode('utf-8')).digest()


def _hash_pair(left, right):
    return hashlib.sha256(left + right).digest()


def _bag_peaks(peaks):
    return peaks[0] if len(peaks) == 1 else hashlib.sha256(b''.join(peaks)).digest()


def build_merkle_tree(leaves):
    """Build an MMR forest from a list of leaf strings.
    Returns (global_root_hash, proofs) where each proof is
    (path, peak_index, other_peak_hashes) and path is [(sibling_hash, is_right)].

    Peaks are perfect binary trees over consecutive leaf ranges whose sizes are
    the set bits of len(leaves), largest first. Proofs are built by leaf
    position, so duplicate chunks each get a correct proof."""
    n = len(leaves)
    if n == 0:
        return None, []

    # Split leaves into perfect-tree peaks (largest first)
    peak_ranges = []
    start = 0
    for bit in reversed(range(n.bit_length())):
        size = 1 << bit
        if n & size:
            peak_ranges.append((start, size))
            start += size

    leaf_hashes = [_hash_leaf(l) for l in leaves]
    peak_levels = []   # per peak: list of levels, level 0 = leaf hashes
    for start, size in peak_ranges:
        levels = [leaf_hashes[start:start + size]]
        while len(levels[-1]) > 1:
            prev = levels[-1]
            levels.append([_hash_pair(prev[i], prev[i + 1]) for i in range(0, len(prev), 2)])
        peak_levels.append(levels)

    peak_hashes = [levels[-1][0] for levels in peak_levels]
    global_root = _bag_peaks(peak_hashes)

    proofs = []
    for peak_idx, (start, size) in enumerate(peak_ranges):
        levels = peak_levels[peak_idx]
        other_peaks = peak_hashes[:peak_idx] + peak_hashes[peak_idx + 1:]
        for offset in range(size):
            path = []
            pos = offset
            for level in levels[:-1]:
                if pos % 2 == 0:
                    path.append((level[pos + 1], True))    # sibling on the right
                else:
                    path.append((level[pos - 1], False))   # sibling on the left
                pos //= 2
            proofs.append((path, peak_idx, other_peaks))
    return global_root, proofs


def verify_merkle_proof(leaf, proof, root_hash):
    """Verify a Merkle proof for a leaf against the root hash."""
    if not proof:
        return False
    path, peak_idx, other_peaks = proof
    current = _hash_leaf(leaf)
    for sibling_hash, is_right in path:
        current = _hash_pair(current, sibling_hash) if is_right else _hash_pair(sibling_hash, current)
    peaks = list(other_peaks)
    peaks.insert(peak_idx, current)
    return _bag_peaks(peaks) == root_hash

# ------------------------------------------------------------------
# Vector normalization (returns a copy)
# ------------------------------------------------------------------
def normalize_vectors(vectors):
    """Normalize vectors to unit length (L2). Returns a new float32 array.
    NOTE: callers must use the return value."""
    vectors = np.asarray(vectors, dtype=np.float32)
    if vectors.ndim == 1:
        n = np.linalg.norm(vectors)
        return vectors / (n if n > 0 else 1e-8)
    norms = np.linalg.norm(vectors, axis=1, keepdims=True)
    norms[norms == 0] = 1e-8
    return np.ascontiguousarray(vectors / norms, dtype=np.float32)

# ------------------------------------------------------------------
# Save / load utilities
# ------------------------------------------------------------------
def file_sha256(path, block=1 << 20):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(block), b''):
            h.update(b)
    return h.hexdigest()


def save_pickle(obj, path):
    with open(path, 'wb') as f:
        pickle.dump(obj, f)


def load_pickle(path):
    with open(path, 'rb') as f:
        return pickle.load(f)
